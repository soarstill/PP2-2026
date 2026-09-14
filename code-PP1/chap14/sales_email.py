import openpyxl, smtplib, sys
from email.mime.text import MIMEText


workbook = openpyxl.load_workbook('구매현황.xlsx')
sheet = workbook.get_sheet_by_name('Sheet1')

members = {}
for r in range(2, sheet.max_row + 1):
    pay = sheet.cell(row=r, column=sheet.max_column).value

    if pay != '구매':
        name = sheet.cell(row=r, column=1).value
        email = sheet.cell(row=r, column=2).value
        members[name] = email


session = smtplib.SMTP_SSL('smtp.naver.com', 465)
session.ehlo()
session.login('abc@naver.com', 'password')


for name, email in members.items():
    msg = MIMEText(f'{name} 회원님, 30% 할인쿠폰이 발행되었습니다. \n 기간 내에 방문해주세요.\n 감사합니다.')
    msg['Subject']='할인쿠폰 증정 행사'
    msg['From']='abc@naver.com'
    msg['To']=email
    print(f'{email}에게 이메일을 보내는 중입니다.')
    status = session.sendmail('abc@naver.com', email, msg.as_string())

    if status != {}:
        print(f'이메일 {email} 전송에서 문제 {status}가 발생하였습니다.' )
session.quit()