from openpyxl import Workbook
workbook = Workbook()

sheet = workbook.active
sheet2 = workbook.create_sheet('Sheet1')
sheet['A1'] = 123

sheet.cell(row=4, column=2, value=10)
workbook.save("test.xlsx")